"""Build a local MOE snapshot from the three officially downloaded workbooks.
Usage: python scripts/build-major-snapshot.py SOURCE_DIR REVIEW_DATE
Requires openpyxl. No network request is made by this script.
"""
import sys, json, re, hashlib
from pathlib import Path
import openpyxl
repo = Path(__file__).resolve().parents[1]
src = Path(sys.argv[1])
date = sys.argv[2]
classes, details = {}, {}
for row in openpyxl.load_workbook(src/'bcode.xlsx', read_only=True, data_only=True).active.values:
    for value in row:
        match = re.fullmatch(r'(\d{4,5})(\D.*)', str(value or '').strip())
        if match:
            (classes if len(match[1]) == 4 else details)[match[1]] = match[2].strip()
assert len(classes) == 93, len(classes)
assert len(details) == 174, len(details)
records = {}
for row in list(openpyxl.load_workbook(src/'students.xlsx', read_only=True, data_only=True).active.values)[3:]:
    school_id, school, code, major, _, level = row[:6]
    code = str(code)
    if not re.fullmatch(r'\d{8}', code): continue
    key = f'u-{school_id}-{code}-{major}'
    records.setdefault(key, dict(id=key, school=school, schoolCode=school_id, major=major, departmentCode=code, classCode=code[:4], detailCode=code[:5], education='專科以上', sourceIds=['moe-students','moe-bcode'], levels=[]))
    if level not in records[key]['levels']: records[key]['levels'].append(level)
workbook = openpyxl.load_workbook(src/'vocational.xlsx', read_only=True, data_only=True)
for row in list(workbook.active.values)[3:]:
    _, _, school_id, school, level, level_name, _, _, group, group_name, code, major = row[:12]
    if level not in ['B','C']: continue
    key = f'v-{school_id}-{code}-{major}'
    records.setdefault(key, dict(id=key, school=school, schoolCode=school_id, major=major, departmentCode=str(code), classCode=str(group), className=group_name, detailCode='', education='高職', sourceIds=['moe-vocational'], levels=[level_name]))
# Preserve the statutory list even where no currently recruiting school appears.
policy = json.loads((repo/'docs/data/hygiene-major-policy.json').read_text(encoding='utf-8'))
for major in policy['vocationalMajors']:
    codes = set()
    for sheet in workbook.worksheets[1:]:
        for row in sheet.values:
            for i, val in enumerate(row):
                if str(val or '').replace('*','').strip() == major and i:
                    code = str(row[i-1])
                    if re.fullmatch(r'[A-Z0-9]{3}',code): codes.add(code)
    key = 'statutory-'+major
    records[key] = dict(id=key, school='各校（法定科別）', schoolCode='', major=major, departmentCode='／'.join(sorted(codes)), classCode='', className='第6條指定科別', detailCode='', education='高職', sourceIds=['hm-law','moe-vocational'], levels=['指定高職科別'], statutory=True)
for record in records.values():
    if record['education'] == '專科以上':
        record['className'] = classes.get(record['classCode'],'未收錄學類')
        record['detailName'] = details.get(record['detailCode'],'未收錄細學類')
    record['lastConfirmed'] = date
result = dict(id='hygiene-majors',name='校系與學類代碼快照',version=date,lastReviewed=date,schoolYear='114（2025–2026）',classificationVersion='第5次修正（106年9月）',sourceIds=['moe-students','moe-vocational','moe-bcode'],coverage='教育部114學年度各校科系別學生數與高級中等學校專業群科／進修部資料；以主要分類碼判讀，未含各系所有相關細學類及歷年改制對照。法定科別另列通用紀錄。',workbooks=[dict(file=f,sha256=hashlib.sha256((src/f).read_bytes()).hexdigest()) for f in ['bcode.xlsx','students.xlsx','vocational.xlsx']],classes=[dict(code=k,name=v) for k,v in sorted(classes.items())],records=list(records.values()))
(repo/'docs/data/hygiene-majors.json').write_text(json.dumps(result,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
print(json.dumps(dict(records=len(records),university=sum(r['education']=='專科以上' for r in records.values()),vocational=sum(r['education']=='高職' for r in records.values())),ensure_ascii=False))
